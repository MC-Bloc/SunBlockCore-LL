'''
SunBlockCore-LL — unified solar monitoring + API server

Written by M. Shahrom Ali (github.com/estineali)
for The SunBlock Project
under the TAG MC-Bloc, Milieux Institute, Concordia University, Montreal, Canada.

https://github.com/MC-Bloc/SunBlock

Setup:
  1. Copy sample.env → .env and fill in values.
  2. Generate a password hash:  python3 scripts/gen_password_hash.py
  3. Vendor frontend assets:    bash scripts/vendor.sh
  4. Run:  uvicorn sunblock:socket_app --host 0.0.0.0 --port ${PORT:-3000}

Simulator mode (no hardware):
  Set SIM_MODE=true in .env — live data is generated from real deployment baselines.

Source layout:
  config.py    — env vars + shared mutable state
  auth.py      — JWT, bcrypt, FastAPI deps, Pydantic models
  db.py        — SQLite helpers + application logging
  hardware.py  — controller polling, power profiles, parameter r/w
  simulator.py — hardware-free data generator
  sunblock.py  — app, lifespan, routes, Socket.IO   ← you are here

Dependencies:
  pip install fastapi uvicorn python-socketio epevermodbus python-dotenv \
              python-jose[cryptography] bcrypt slowapi jinja2
'''

import asyncio
import csv
import io
import os
from contextlib import asynccontextmanager
from datetime import datetime
from typing import Optional

import bcrypt as _bcrypt
import socketio
from epevermodbus.driver import EpeverChargeController
from fastapi import Cookie, Depends, FastAPI, HTTPException, Query, Request, Response
from fastapi.exceptions import RequestValidationError
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from slowapi import _rate_limit_exceeded_handler
from slowapi.errors import RateLimitExceeded

import config
import pyotp

from auth import (
    ControllerParamsUpdate, LoginRequest, SettingsUpdate, PasswordChange,
    TokenCreateRequest, TwoFACodeRequest, TwoFADisableRequest, limiter,
    PENDING_2FA_MINUTES,
    check_session, create_token, create_pending_2fa_token, verify_pending_2fa_token,
    verify_totp_code, verify_session, verify_session_or_token,
    require_controller, require_real_controller,
)
from db import (
    admin_log, apply_data_directory, check_db, clear_backup_codes,
    count_unused_backup_codes, create_api_token, delete_setting,
    generate_backup_codes, list_api_tokens, load_settings, query_history,
    query_visualize, revoke_api_token, save_setting, sunblock_log,
    verify_and_consume_backup_code, write_db,
)
from db import VIZ_FIELD_META
from hardware import (
    apply_controller_params, check_power_profile,
    parse_data, read_controller_params,
    read_controller_stats, read_controller_status,
    set_power_profile,
)
from simulator import simulate_data


# ── Polling loop ──────────────────────────────────────────────────────────────

async def polling_loop():
    await sunblock_log("Waking Up...")
    if config.DATA_MAN:
        await sunblock_log("Data Management is " + str(config.DATA_MAN))
        await asyncio.to_thread(check_db)

    loop = asyncio.get_running_loop()

    while config.POLLING_ACTIVE:
        tick_start = loop.time()

        poll_fn = simulate_data if config.SIM_MODE else parse_data
        try:
            new_data = await loop.run_in_executor(None, poll_fn)
            config.SOLAR_DATA = new_data  # atomic reference swap
        except Exception as e:
            await sunblock_log("Hardware error, stopping poll: " + str(e))
            break

        if config.DATA_MAN and not config.SIM_MODE:
            try:
                await loop.run_in_executor(None, write_db)
            except Exception as e:
                await sunblock_log("DB write error (continuing): " + str(e))

        try:
            await sio.emit("solar_data", {**config.SOLAR_DATA, "ConnectedUsers": config.ACTIVE_USERS})
        except Exception as e:
            await sunblock_log("Socket emit error (continuing): " + str(e))

        # Sleep only for the remainder of the interval so the total cycle time
        # stays at READ_INTERVAL regardless of how long the poll and DB write took.
        elapsed = loop.time() - tick_start
        await asyncio.sleep(max(0.0, config.READ_INTERVAL - elapsed))

    await sunblock_log("Exiting polling loop.")


# ── Lifespan ──────────────────────────────────────────────────────────────────

@asynccontextmanager
async def lifespan(app: FastAPI):
    await asyncio.to_thread(load_settings)
    await sunblock_log("Settings loaded from persistent store.")

    if config.DATA_DIRECTORY:
        os.makedirs(config.DATA_DIRECTORY, exist_ok=True)
    else:
        await sunblock_log(
            "WARNING: DATA_DIRECTORY is not set. "
            "Open the admin panel → Settings to configure it."
        )

    if not config.ADMIN_PASSWORD_HASH:
        await sunblock_log(
            "WARNING: ADMIN_PASSWORD_HASH not set. "
            "Run scripts/gen_password_hash.py and set it in .env. Login is disabled."
        )

    if not config.ADMIN_PATH:
        await sunblock_log(
            "WARNING: ADMIN_PATH is not set. "
            "The admin login page is unreachable. Set ADMIN_PATH to a secret slug in .env "
            "(e.g. ADMIN_PATH=xK9mP3qR7) then restart."
        )
    else:
        await sunblock_log("Admin route registered (path configured in ADMIN_PATH).")

    if config.SIM_MODE:
        await sunblock_log("SIM_MODE=true — skipping hardware init, using simulator.")
    else:
        try:
            config.CONTROLLER = EpeverChargeController(config.CONTROLLER_PORT, config.CONTROLLER_SLAVE)
        except Exception as e:
            await sunblock_log("Failed to connect to controller: " + str(e))

    config.ACTIVE_USERS_LOCK = asyncio.Lock()

    if config.CONTROLLER is not None or config.SIM_MODE:
        if config.DATA_DIRECTORY:
            config.POLLING_ACTIVE = True
            config.POLLING_TASK = asyncio.create_task(polling_loop())
        else:
            await sunblock_log(
                "Polling deferred — DATA_DIRECTORY not set. "
                "Configure it in the admin panel to begin collecting data."
            )
    else:
        await sunblock_log("Controller unavailable — polling disabled.")

    yield

    if config.POLLING_TASK is not None:
        config.POLLING_ACTIVE = False
        config.POLLING_TASK.cancel()
        try:
            await config.POLLING_TASK
        except asyncio.CancelledError:
            pass
    if config.DB_CONNECTION:
        config.DB_CONNECTION.close()
    await sunblock_log("Server shutting down.")


# ── App setup ─────────────────────────────────────────────────────────────────

sio       = socketio.AsyncServer(async_mode="asgi", cors_allowed_origins=[])
app       = FastAPI(lifespan=lifespan, docs_url=None, redoc_url=None, openapi_url=None)
templates = Jinja2Templates(directory="templates")

app.state.limiter = limiter

# Holds a freshly generated TOTP secret between POST /api/2fa/setup and
# POST /api/2fa/confirm. Deliberately kept in memory only (never persisted) —
# an unconfirmed secret must not be able to silently activate 2FA (e.g. via a
# server restart resuming a half-finished enrollment), and it must not survive
# a restart that the admin didn't initiate. Single-admin model: one pending
# enrollment at a time is sufficient. Cleared on confirm, on a fresh /setup
# call, and never written to disk.
_pending_totp_secret: Optional[str] = None
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)
app.mount("/static", StaticFiles(directory="public"), name="static")

socket_app = socketio.ASGIApp(sio, app)


# ── Security headers ──────────────────────────────────────────────────────────

@app.middleware("http")
async def security_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers["X-Content-Type-Options"]  = "nosniff"
    response.headers["X-Frame-Options"]          = "DENY"
    response.headers["Referrer-Policy"]          = "strict-origin-when-cross-origin"
    response.headers["X-XSS-Protection"]         = "1; mode=block"
    response.headers["Permissions-Policy"]       = "geolocation=(), microphone=(), camera=()"
    return response


# ── Error handlers ────────────────────────────────────────────────────────────

@app.exception_handler(404)
async def not_found(request: Request, exc: HTTPException):
    # Return JSON for API paths, HTML for everything else
    if request.url.path.startswith("/api/") or request.url.path.startswith("/socket.io/"):
        return JSONResponse(status_code=404, content={"detail": "Not found"})
    return templates.TemplateResponse("404.html", {"request": request}, status_code=404)


# ── Helpers ───────────────────────────────────────────────────────────────────

def _page_response(template: str, context: dict) -> "TemplateResponse":
    """
    Render an HTML page with a CSP header attached.

    All of our JS now lives in same-origin static files (vendor/*, js/sunblock.js)
    loaded via <script src="..."> — there is no inline <script> left in any
    template, so 'self' alone covers script loading and we don't need a
    per-request nonce or 'unsafe-inline' for scripts. (Previously this function
    minted a nonce to allow exactly one inline <script> block; that block was
    extracted to public/js/sunblock.js, so the nonce machinery is obsolete.)

    We still need 'unsafe-eval' because Alpine.js evaluates x-data/x-on
    expressions with new Function() internally, and 'unsafe-inline' for
    style-src because Alpine's :style bindings set inline style="" attributes.
    """
    csp = (
        f"default-src 'self'; "
        f"script-src 'self' 'unsafe-eval'; "
        f"style-src 'self' 'unsafe-inline'; "
        f"img-src 'self' data: blob:; "
        f"connect-src 'self'; "
        f"frame-ancestors 'none'"
    )
    response = templates.TemplateResponse(template, context)
    response.headers["Content-Security-Policy"] = csp
    return response


# ── Routes ────────────────────────────────────────────────────────────────────

@app.get("/")
async def index(request: Request):
    return _page_response("public.html", {
        "request":          request,
        "is_authenticated": check_session(request),
        "sim_mode":         config.SIM_MODE,
        "env_defaults":     config.ENV_DEFAULTS,
        "data_directory":   config.DATA_DIRECTORY or "",
        "admin_mode":       False,
        "viz_fields":       VIZ_FIELD_META,
    })

async def _admin_page(request: Request):
    """Full admin panel — all tabs and modals rendered server-side (admin_mode=True).
    Registered at the secret ADMIN_PATH slug; not exposed at any predictable URL.
    Templates: admin.html extends _base.html and fills in every admin-only block —
    see templates/_base.html / templates/admin.html / templates/public.html."""
    return _page_response("admin.html", {
        "request":          request,
        "is_authenticated": check_session(request),
        "sim_mode":         config.SIM_MODE,
        "env_defaults":     config.ENV_DEFAULTS,
        "data_directory":   config.DATA_DIRECTORY or "",
        "admin_mode":       True,
        "viz_fields":       VIZ_FIELD_META,
    })

# Register the admin route only if ADMIN_PATH is configured.
if config.ADMIN_PATH:
    _path = "/" + config.ADMIN_PATH.lstrip("/")
    app.add_api_route(_path, _admin_page, methods=["GET"], include_in_schema=False)

@app.get("/api/mode")
async def get_mode():
    return {"mode": "simulator" if config.SIM_MODE else "live"}


# Auth

@app.post("/api/login")
@limiter.limit("5/minute")
async def login(request: Request, body: LoginRequest, response: Response):
    ip = _client_ip(request)
    if not config.ADMIN_PASSWORD_HASH:
        raise HTTPException(status_code=503, detail="Admin password not configured on server.")
    if body.username != config.ADMIN_USERNAME or \
       not _bcrypt.checkpw(body.password.encode(), config.ADMIN_PASSWORD_HASH.encode()):
        await admin_log("LOGIN_FAILED", f"user={body.username}", ip=ip)
        raise HTTPException(status_code=401, detail="Invalid credentials")

    if config.TOTP_ENABLED:
        # Password alone is not enough — issue a short-lived "pending" token in
        # its own cookie (never `sb_session`) and require a second factor before
        # any session is granted. See auth.create_pending_2fa_token.
        response.set_cookie(
            key="sb_2fa_pending",
            value=create_pending_2fa_token(body.username),
            httponly=True,
            secure=config.SECURE_COOKIES,
            samesite="strict",
            max_age=PENDING_2FA_MINUTES * 60,
        )
        await admin_log("LOGIN_PASSWORD_OK_2FA_PENDING", f"user={body.username}", ip=ip)
        return {"requires_2fa": True, "message": "Password verified — enter your 2FA code"}

    response.set_cookie(
        key="sb_session",
        value=create_token(body.username),
        httponly=True,
        secure=config.SECURE_COOKIES,
        samesite="strict",
        max_age=config.TOKEN_EXPIRE_HOURS * 3600,
    )
    await admin_log("LOGIN", f"user={body.username}", ip=ip)
    return {"message": "Logged in"}


@app.post("/api/login/verify-2fa")
@limiter.limit("5/minute")
async def login_verify_2fa(request: Request, body: TwoFACodeRequest, response: Response,
                           sb_2fa_pending: Optional[str] = Cookie(default=None)):
    ip = _client_ip(request)
    username = verify_pending_2fa_token(sb_2fa_pending)
    if not username:
        raise HTTPException(status_code=401, detail="2FA challenge expired or invalid — log in again")

    code = body.code.strip()
    ok = verify_totp_code(code)
    used_backup = False
    if not ok:
        ok = await asyncio.to_thread(verify_and_consume_backup_code, code)
        used_backup = ok

    if not ok:
        await admin_log("2FA_CHALLENGE_FAILED", f"user={username}", ip=ip)
        raise HTTPException(status_code=401, detail="Invalid authentication code")

    response.delete_cookie("sb_2fa_pending", samesite="strict")
    response.set_cookie(
        key="sb_session",
        value=create_token(username),
        httponly=True,
        secure=config.SECURE_COOKIES,
        samesite="strict",
        max_age=config.TOKEN_EXPIRE_HOURS * 3600,
    )
    if used_backup:
        await admin_log("BACKUP_CODE_USED", f"user={username}", ip=ip)
        remaining = await asyncio.to_thread(count_unused_backup_codes)
        if remaining <= 2:
            await sunblock_log(
                f"WARNING: only {remaining} unused 2FA backup codes remain for "
                f"user={username}. Generate new ones from Settings soon."
            )
    await admin_log("2FA_LOGIN", f"user={username} method={'backup_code' if used_backup else 'totp'}", ip=ip)
    return {"message": "Logged in"}


@app.post("/api/logout")
async def logout(request: Request, response: Response):
    await admin_log("LOGOUT", ip=_client_ip(request))
    response.delete_cookie("sb_session", samesite="strict")
    response.delete_cookie("sb_2fa_pending", samesite="strict")
    return {"message": "Logged out"}

@app.get("/api/auth/status")
async def auth_status(request: Request):
    return {"authenticated": check_session(request)}


# Live data

@app.get("/api/data")
async def get_data():
    return JSONResponse(content=config.SOLAR_DATA, status_code=200)


@app.get("/api/data/history")
@limiter.limit("60/minute")
async def get_data_history(
    request: Request,
    limit:   int                = 100,
    offset:  int                = 0,
    from_ts: Optional[str]      = Query(default=None, alias="from"),
    to_ts:   Optional[str]      = Query(default=None, alias="to"),
    order:   str                = "desc",
    user:    str                = Depends(verify_session_or_token),
):
    """
    Paginated history of solar readings from the SQLite database.

    Query params:
      limit   — rows per page (1–1000, default 100)
      offset  — row index to start from (default 0)
      from    — optional ISO date/datetime lower bound (e.g. 2025-05-11)
      to      — optional ISO date/datetime upper bound (e.g. 2025-05-15)
      order   — "desc" (newest first, default) or "asc" (oldest first)

    Returns {"rows": [...], "total": N, "limit": N, "offset": N}.
    Rate-limited to 60 requests/minute per IP to prevent DoS on large DBs.
    """
    loop = asyncio.get_running_loop()
    return await loop.run_in_executor(
        None, query_history, limit, offset, from_ts, to_ts, order
    )


def _client_ip(request: Request) -> str:
    """Return the connecting client's IP address, or 'unknown'."""
    return request.client.host if request.client else "unknown"


def _db_guard():
    """Raise 404 if the telemetry DB doesn't exist yet."""
    if not config.DB_NAME or not os.path.isfile(config.DB_NAME):
        raise HTTPException(
            status_code=404,
            detail="Database file not found. DATA_DIRECTORY may not be configured or data collection may be disabled.",
        )

@app.get("/api/data/visualize/fields")
@limiter.limit("60/minute")
async def get_viz_fields(request: Request, user: str = Depends(verify_session_or_token)):
    """Return the list of plottable fields with label and unit metadata."""
    return VIZ_FIELD_META

@app.get("/api/data/visualize")
@limiter.limit("20/minute")
async def get_visualize(
    request:       Request,
    fields:        str           = Query(default="BattPercentage,PVPower,LoadPower"),
    from_ts:       Optional[str] = Query(default=None, alias="from"),
    to_ts:         Optional[str] = Query(default=None, alias="to"),
    sample:        int           = Query(default=1,  ge=1, le=3600),
    smooth:        int           = Query(default=0,  ge=0, le=300),
    filter_spikes: bool          = Query(default=True),
    user:          str           = Depends(verify_session_or_token),
):
    """
    Time-series data for the Visualize tab.

    Processing: date filter → row resampling → spike filter → smoothing.
    Returns timestamps + one series object per requested field.
    Rate-limited to 20 requests/minute per IP to prevent DoS on large DBs.
    """
    field_list = [f.strip() for f in fields.split(",") if f.strip()]
    loop = asyncio.get_running_loop()
    return await loop.run_in_executor(
        None, query_visualize, field_list, from_ts, to_ts, sample, smooth, filter_spikes
    )

@app.get("/api/data/download")
async def download_db(request: Request, user: str = Depends(verify_session_or_token)):
    """Download the raw SQLite telemetry database file (auth required)."""
    _db_guard()
    await admin_log("DOWNLOAD", "format=sqlite", ip=_client_ip(request))
    return FileResponse(
        path=config.DB_NAME,
        media_type="application/octet-stream",
        filename="SunBlockCore-LL.db",
    )

@app.get("/api/data/download/csv")
async def download_csv(request: Request, user: str = Depends(verify_session_or_token)):
    """Download all telemetry rows as a CSV file (auth required)."""
    _db_guard()
    await admin_log("DOWNLOAD", "format=csv", ip=_client_ip(request))

    def _build():
        result = query_history(limit=1_000_000, offset=0, from_ts=None, to_ts=None, order="asc")
        rows = result["rows"]
        buf = io.StringIO()
        if rows:
            writer = csv.DictWriter(buf, fieldnames=rows[0].keys())
            writer.writeheader()
            writer.writerows(rows)
        return buf.getvalue()

    loop = asyncio.get_running_loop()
    content = await loop.run_in_executor(None, _build)
    return StreamingResponse(
        iter([content]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=SunBlockCore-LL.csv"},
    )

@app.get("/api/data/download/xlsx")
async def download_xlsx(request: Request, user: str = Depends(verify_session_or_token)):
    """Download all telemetry rows as an Excel file (auth required)."""
    _db_guard()
    await admin_log("DOWNLOAD", "format=xlsx", ip=_client_ip(request))

    def _build():
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        result = query_history(limit=1_000_000, offset=0, from_ts=None, to_ts=None, order="asc")
        rows = result["rows"]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Solar Data"

        if rows:
            headers = list(rows[0].keys())
            # Header row styling
            header_fill = PatternFill("solid", fgColor="1a1d27")
            header_font = Font(bold=True, color="F59E0B")
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            # Data rows
            for row_idx, row in enumerate(rows, 2):
                for col_idx, key in enumerate(headers, 1):
                    ws.cell(row=row_idx, column=col_idx, value=row[key])

            # Auto-fit column widths
            for col in ws.columns:
                max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 30)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    loop = asyncio.get_running_loop()
    content = await loop.run_in_executor(None, _build)
    return StreamingResponse(
        iter([content]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=SunBlockCore-LL.xlsx"},
    )


# Settings

def _settings_snapshot() -> dict:
    return {
        "data_directory":     config.DATA_DIRECTORY or "",
        "read_interval":      config.READ_INTERVAL,
        "data_man":           config.DATA_MAN,
        "sim_mode":           config.SIM_MODE,
        "token_expire_hours": config.TOKEN_EXPIRE_HOURS,
    }

@app.get("/api/settings")
async def get_settings(user: str = Depends(verify_session_or_token)):
    return _settings_snapshot()

@app.patch("/api/settings")
async def update_settings(request: Request, body: SettingsUpdate, user: str = Depends(verify_session_or_token)):
    ip = _client_ip(request)
    changed: list = []

    if body.data_directory is not None:
        d = body.data_directory.strip()
        if not d:
            raise HTTPException(status_code=400, detail="data_directory cannot be empty")
        try:
            await asyncio.to_thread(apply_data_directory, d)
        except ValueError as e:
            await admin_log("PATH_REJECTED", f"attempted={d}  reason={e}", ip=ip)
            raise HTTPException(status_code=400, detail=str(e))
        except OSError as e:
            raise HTTPException(status_code=400, detail=f"Cannot create directory: {e}")
        await asyncio.to_thread(save_setting, "data_directory", config.DATA_DIRECTORY)
        changed.append(f"data_directory={config.DATA_DIRECTORY}")
        # Initialise telemetry DB now that the directory exists.
        if config.DATA_MAN:
            await asyncio.to_thread(check_db)
        # Start polling if it was held back only because DATA_DIRECTORY was missing.
        if not config.POLLING_ACTIVE and (config.CONTROLLER is not None or config.SIM_MODE):
            config.POLLING_ACTIVE = True
            config.POLLING_TASK = asyncio.create_task(polling_loop())

    if body.read_interval is not None:
        if not 1 <= body.read_interval <= 3600:
            raise HTTPException(status_code=400, detail="read_interval must be 1–3600")
        config.READ_INTERVAL = body.read_interval
        await asyncio.to_thread(save_setting, "read_interval", body.read_interval)
        changed.append(f"read_interval={body.read_interval}")

    if body.data_man is not None:
        config.DATA_MAN = body.data_man
        await asyncio.to_thread(save_setting, "data_man", body.data_man)
        changed.append(f"data_man={body.data_man}")

    if body.sim_mode is not None:
        if not body.sim_mode and config.CONTROLLER is None:
            raise HTTPException(status_code=400, detail="Cannot disable simulator — no hardware controller is connected")
        config.SIM_MODE = body.sim_mode
        await asyncio.to_thread(save_setting, "sim_mode", body.sim_mode)
        changed.append(f"sim_mode={body.sim_mode}")

    if body.token_expire_hours is not None:
        if not 1 <= body.token_expire_hours <= 720:
            raise HTTPException(status_code=400, detail="token_expire_hours must be 1–720")
        config.TOKEN_EXPIRE_HOURS = body.token_expire_hours
        await asyncio.to_thread(save_setting, "token_expire_hours", body.token_expire_hours)
        changed.append(f"token_expire_hours={body.token_expire_hours}")

    if changed:
        await admin_log("SETTINGS_CHANGE", "  ".join(changed), ip=ip)

    return _settings_snapshot()

@app.delete("/api/settings/{key}")
async def reset_setting(key: str, request: Request, user: str = Depends(verify_session_or_token)):
    _restorable = {"read_interval", "data_man", "sim_mode", "token_expire_hours"}
    if key not in _restorable:
        raise HTTPException(status_code=404, detail=f"Unknown or non-resettable setting: {key}")

    env_val = config.ENV_DEFAULTS[key]

    if key == "sim_mode" and not env_val and config.CONTROLLER is None:
        raise HTTPException(status_code=400, detail="Cannot disable simulator — no hardware controller is connected")

    await asyncio.to_thread(delete_setting, key)

    # Restore live config to the original env value
    _restore = {
        "read_interval":      lambda v: setattr(config, "READ_INTERVAL",      v),
        "data_man":           lambda v: setattr(config, "DATA_MAN",           v),
        "sim_mode":           lambda v: setattr(config, "SIM_MODE",           v),
        "token_expire_hours": lambda v: setattr(config, "TOKEN_EXPIRE_HOURS", v),
    }
    _restore[key](env_val)

    await admin_log("RESET_SETTING", f"key={key}  reverted_to={env_val}", ip=_client_ip(request))
    return _settings_snapshot()


@app.post("/api/settings/password")
async def change_password(request: Request, body: PasswordChange, user: str = Depends(verify_session)):
    ip = _client_ip(request)
    if not config.ADMIN_PASSWORD_HASH:
        raise HTTPException(status_code=503, detail="No password configured on server")
    if not _bcrypt.checkpw(body.current_password.encode(), config.ADMIN_PASSWORD_HASH.encode()):
        await admin_log("PASSWORD_CHANGE_FAILED", f"user={user}", ip=ip)
        raise HTTPException(status_code=400, detail="Current password is incorrect")
    if len(body.new_password) < 8:
        raise HTTPException(status_code=400, detail="New password must be at least 8 characters")
    new_hash = _bcrypt.hashpw(body.new_password.encode(), _bcrypt.gensalt()).decode()
    config.ADMIN_PASSWORD_HASH = new_hash
    await asyncio.to_thread(save_setting, "admin_password_hash", new_hash)
    await admin_log("PASSWORD_CHANGE", f"user={user}", ip=ip)
    return {"message": "Password updated"}


# Two-factor authentication (TOTP) — enroll / confirm / disable.
#
# All of these are session-only (Depends(verify_session)), never bearer-token
# eligible — exactly like change_password and token management. A leaked API
# token must never be able to touch the account's second factor.
#
# Enrollment is two steps (setup → confirm) so a typo'd authenticator app never
# locks the admin out: the secret only becomes active (persisted + totp_enabled
# set) after the admin proves they can already generate valid codes with it.

@app.get("/api/2fa/status")
async def twofa_status(user: str = Depends(verify_session)):
    return {
        "enabled": config.TOTP_ENABLED,
        "backup_codes_remaining": await asyncio.to_thread(count_unused_backup_codes) if config.TOTP_ENABLED else 0,
    }


@app.post("/api/2fa/setup")
async def twofa_setup(request: Request, user: str = Depends(verify_session)):
    global _pending_totp_secret
    if config.TOTP_ENABLED:
        raise HTTPException(status_code=400, detail="2FA is already enabled — disable it first to re-enroll")

    _pending_totp_secret = pyotp.random_base32()
    uri = pyotp.totp.TOTP(_pending_totp_secret).provisioning_uri(
        name=config.ADMIN_USERNAME, issuer_name="SunBlockCore-LL"
    )
    await admin_log("2FA_SETUP_STARTED", f"user={user}", ip=_client_ip(request))
    return {
        "secret": _pending_totp_secret,
        "otpauth_uri": uri,
        "message": "Scan the QR code (or enter the secret manually) in your authenticator app, "
                   "then submit the 6-digit code it generates to confirm.",
    }


@app.post("/api/2fa/confirm")
async def twofa_confirm(request: Request, body: TwoFACodeRequest, user: str = Depends(verify_session)):
    global _pending_totp_secret
    ip = _client_ip(request)
    if config.TOTP_ENABLED:
        raise HTTPException(status_code=400, detail="2FA is already enabled")
    if not _pending_totp_secret:
        raise HTTPException(status_code=400, detail="No 2FA setup in progress — call /api/2fa/setup first")

    code = body.code.strip().replace(" ", "")
    try:
        valid = pyotp.TOTP(_pending_totp_secret).verify(code, valid_window=1)
    except Exception:
        valid = False
    if not valid:
        await admin_log("2FA_CHALLENGE_FAILED", f"user={user} context=confirm", ip=ip)
        raise HTTPException(status_code=400, detail="Incorrect code — check your authenticator app and try again")

    secret = _pending_totp_secret
    _pending_totp_secret = None
    config.TOTP_SECRET = secret
    config.TOTP_ENABLED = True
    await asyncio.to_thread(save_setting, "totp_secret", secret)
    await asyncio.to_thread(save_setting, "totp_enabled", True)
    backup_codes = await asyncio.to_thread(generate_backup_codes)
    await admin_log("2FA_ENABLED", f"user={user}", ip=ip)
    return {
        "message": "Two-factor authentication is now enabled.",
        "backup_codes": backup_codes,
        "backup_codes_notice": "Save these somewhere safe — they will not be shown again. "
                               "Each one can be used once to log in if you lose access to your authenticator app.",
    }


@app.post("/api/2fa/disable")
async def twofa_disable(request: Request, body: TwoFADisableRequest, user: str = Depends(verify_session)):
    global _pending_totp_secret
    ip = _client_ip(request)
    if not config.TOTP_ENABLED:
        raise HTTPException(status_code=400, detail="2FA is not enabled")

    # Defense in depth — mirrors change_password: require BOTH the current
    # password AND a valid second-factor code to turn 2FA off. Otherwise a
    # hijacked session alone (e.g. via XSS) could strip the account's strongest
    # protection with one click.
    if not _bcrypt.checkpw(body.password.encode(), config.ADMIN_PASSWORD_HASH.encode()):
        await admin_log("2FA_DISABLE_FAILED", f"user={user} reason=bad_password", ip=ip)
        raise HTTPException(status_code=400, detail="Current password is incorrect")

    code = body.code.strip()
    ok = verify_totp_code(code) or await asyncio.to_thread(verify_and_consume_backup_code, code)
    if not ok:
        await admin_log("2FA_DISABLE_FAILED", f"user={user} reason=bad_code", ip=ip)
        raise HTTPException(status_code=400, detail="Incorrect authentication code")

    config.TOTP_ENABLED = False
    config.TOTP_SECRET = ""
    _pending_totp_secret = None
    await asyncio.to_thread(save_setting, "totp_enabled", False)
    await asyncio.to_thread(delete_setting, "totp_secret")
    await asyncio.to_thread(clear_backup_codes)
    await admin_log("2FA_DISABLED", f"user={user}", ip=ip)
    return {"message": "Two-factor authentication has been disabled."}


@app.post("/api/2fa/backup-codes/regenerate")
async def twofa_regenerate_backup_codes(request: Request, body: TwoFACodeRequest, user: str = Depends(verify_session)):
    """Invalidate all existing backup codes and issue a fresh set — requires a valid 2FA code."""
    ip = _client_ip(request)
    if not config.TOTP_ENABLED:
        raise HTTPException(status_code=400, detail="2FA is not enabled")
    code = body.code.strip()
    ok = verify_totp_code(code) or await asyncio.to_thread(verify_and_consume_backup_code, code)
    if not ok:
        await admin_log("2FA_CHALLENGE_FAILED", f"user={user} context=regenerate_backup_codes", ip=ip)
        raise HTTPException(status_code=400, detail="Incorrect authentication code")

    backup_codes = await asyncio.to_thread(generate_backup_codes)
    await admin_log("BACKUP_CODES_REGENERATED", f"user={user}", ip=ip)
    return {
        "backup_codes": backup_codes,
        "backup_codes_notice": "Save these somewhere safe — they will not be shown again. "
                               "All previous backup codes have been invalidated.",
    }


# API tokens — for authenticating to the API from outside the browser.
#
# Management endpoints (create/list/revoke) intentionally require a real
# session cookie, NOT a bearer token — otherwise a leaked token could be used
# to mint further tokens or revoke the admin's own, escalating a single leak
# into permanent persistence. A token is good for data/control-plane access
# only, never for managing the token store itself.

@app.post("/api/tokens")
async def create_token_route(request: Request, body: TokenCreateRequest, user: str = Depends(verify_session)):
    ip = _client_ip(request)
    name = body.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Token name is required")
    if body.expires_in_hours is not None and not (1 <= body.expires_in_hours <= 8760):
        raise HTTPException(status_code=400, detail="expires_in_hours must be 1–8760 (or omitted for no expiry)")

    token_id, raw_token = await asyncio.to_thread(create_api_token, name, body.expires_in_hours)
    await admin_log(
        "TOKEN_CREATED",
        f"id={token_id} name={name} expires_in_hours={body.expires_in_hours or 'never'}",
        ip=ip,
    )
    return {
        "id": token_id,
        "name": name,
        "token": raw_token,
        "message": "Save this token now — it will not be shown again.",
    }


@app.get("/api/tokens")
async def list_tokens_route(user: str = Depends(verify_session)):
    return {"tokens": await asyncio.to_thread(list_api_tokens)}


@app.delete("/api/tokens/{token_id}")
async def revoke_token_route(token_id: int, request: Request, user: str = Depends(verify_session)):
    revoked = await asyncio.to_thread(revoke_api_token, token_id)
    if not revoked:
        raise HTTPException(status_code=404, detail="Token not found")
    await admin_log("TOKEN_REVOKED", f"id={token_id}", ip=_client_ip(request))
    return {"message": "Token revoked"}


# Power profile

@app.get("/api/power-profile")
async def get_power_profile(_=Depends(require_controller)):
    loop = asyncio.get_running_loop()
    return {"profile": await loop.run_in_executor(None, check_power_profile)}

@app.post("/api/performance-mode")
async def set_performance(request: Request, user: str = Depends(verify_session_or_token), _=Depends(require_real_controller)):
    loop = asyncio.get_running_loop()
    result = await loop.run_in_executor(None, set_power_profile, "performance")
    await admin_log("POWER_PROFILE", "profile=performance", ip=_client_ip(request))
    return {"profile": result}

@app.post("/api/power-saver-mode")
async def set_power_saver(request: Request, user: str = Depends(verify_session_or_token), _=Depends(require_real_controller)):
    loop = asyncio.get_running_loop()
    result = await loop.run_in_executor(None, set_power_profile, "power-saver")
    await admin_log("POWER_PROFILE", "profile=power-saver", ip=_client_ip(request))
    return {"profile": result}

@app.post("/api/balanced")
async def set_balanced(request: Request, user: str = Depends(verify_session_or_token), _=Depends(require_real_controller)):
    loop = asyncio.get_running_loop()
    result = await loop.run_in_executor(None, set_power_profile, "balanced")
    await admin_log("POWER_PROFILE", "profile=balanced", ip=_client_ip(request))
    return {"profile": result}


# Controller parameters

@app.get("/api/controller/parameters")
async def get_controller_params(_=Depends(require_controller)):
    loop = asyncio.get_running_loop()
    return await loop.run_in_executor(None, read_controller_params)

@app.put("/api/controller/parameters")
async def update_controller_params(
    request: Request,
    body: ControllerParamsUpdate,
    user: str = Depends(verify_session_or_token),
    _=Depends(require_real_controller),
):
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(None, apply_controller_params, body.model_dump())
    await admin_log("CONTROLLER_PARAMS_UPDATE", ip=_client_ip(request))
    return await loop.run_in_executor(None, read_controller_params)

@app.get("/api/controller/stats")
async def get_controller_stats(_=Depends(require_controller)):
    loop = asyncio.get_running_loop()
    return await loop.run_in_executor(None, read_controller_stats)

@app.get("/api/controller/status")
async def get_controller_status(_=Depends(require_controller)):
    loop = asyncio.get_running_loop()
    return await loop.run_in_executor(None, read_controller_status)

@app.post("/api/controller/rtc/sync")
async def sync_rtc(request: Request, user: str = Depends(verify_session_or_token), _=Depends(require_real_controller)):
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(None, lambda: config.CONTROLLER.set_rtc(datetime.now()))
    await admin_log("RTC_SYNC", ip=_client_ip(request))
    return {"message": "RTC synced to server time"}


# ── Socket.IO events ──────────────────────────────────────────────────────────

@sio.event
async def connect(sid, environ):
    async with config.ACTIVE_USERS_LOCK:
        config.ACTIVE_USERS += 1
    print("Client connected:", sid)

@sio.event
async def disconnect(sid):
    async with config.ACTIVE_USERS_LOCK:
        config.ACTIVE_USERS -= 1
    print("Client disconnected:", sid)
